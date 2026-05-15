"""Генерация документов (.docx и .pdf) для AI-агентства."""

import os
from datetime import datetime

from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import cm
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont


OUTPUT_DIR = "/tmp"


async def generate_docx(
    order_id: int,
    service_name: str,
    result_text: str,
    date: str = None,
) -> str:
    """
    Генерация .docx файла с результатом заказа.

    Args:
        order_id: ID заказа
        service_name: название услуги
        result_text: текст результата
        date: дата выполнения (если None - текущая)

    Returns:
        Путь к сгенерированному файлу
    """
    if date is None:
        date = datetime.utcnow().strftime("%Y-%m-%d %H:%M")

    doc = Document()

    # Заголовок
    title = doc.add_heading("AI-Agency", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Метаданные
    doc.add_paragraph(f"Order #{order_id}")
    doc.add_paragraph(f"Service: {service_name}")
    doc.add_paragraph(f"Date: {date}")
    doc.add_paragraph("")  # пустая строка-разделитель

    # Контент
    doc.add_heading("Result", level=1)
    for paragraph_text in result_text.split("\n"):
        if paragraph_text.strip():
            doc.add_paragraph(paragraph_text)

    # Сохранение
    filepath = os.path.join(OUTPUT_DIR, f"order_{order_id}.docx")
    doc.save(filepath)
    return filepath


async def generate_pdf(
    order_id: int,
    service_name: str,
    result_text: str,
    date: str = None,
) -> str:
    """
    Генерация .pdf файла с результатом заказа.

    Args:
        order_id: ID заказа
        service_name: название услуги
        result_text: текст результата
        date: дата выполнения (если None - текущая)

    Returns:
        Путь к сгенерированному файлу
    """
    if date is None:
        date = datetime.utcnow().strftime("%Y-%m-%d %H:%M")

    filepath = os.path.join(OUTPUT_DIR, f"order_{order_id}.pdf")

    doc = SimpleDocTemplate(
        filepath,
        pagesize=A4,
        rightMargin=2 * cm,
        leftMargin=2 * cm,
        topMargin=2 * cm,
        bottomMargin=2 * cm,
    )

    styles = getSampleStyleSheet()

    # Стиль заголовка
    title_style = ParagraphStyle(
        "CustomTitle",
        parent=styles["Title"],
        fontSize=24,
        spaceAfter=20,
    )

    # Стиль метаданных
    meta_style = ParagraphStyle(
        "Meta",
        parent=styles["Normal"],
        fontSize=10,
        spaceAfter=4,
    )

    # Стиль контента
    content_style = ParagraphStyle(
        "Content",
        parent=styles["Normal"],
        fontSize=11,
        spaceAfter=8,
        leading=14,
    )

    elements = []

    # Заголовок
    elements.append(Paragraph("AI-Agency", title_style))
    elements.append(Spacer(1, 12))

    # Метаданные
    elements.append(Paragraph(f"Order #{order_id}", meta_style))
    elements.append(Paragraph(f"Service: {service_name}", meta_style))
    elements.append(Paragraph(f"Date: {date}", meta_style))
    elements.append(Spacer(1, 20))

    # Контент
    elements.append(Paragraph("Result", styles["Heading2"]))
    elements.append(Spacer(1, 8))

    for paragraph_text in result_text.split("\n"):
        if paragraph_text.strip():
            # Экранируем спецсимволы для reportlab
            safe_text = (
                paragraph_text
                .replace("&", "&amp;")
                .replace("<", "&lt;")
                .replace(">", "&gt;")
            )
            elements.append(Paragraph(safe_text, content_style))

    doc.build(elements)
    return filepath
