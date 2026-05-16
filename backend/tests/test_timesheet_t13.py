"""Tests for T-13 timesheet export endpoint."""

import io

import pytest
import pytest_asyncio
from httpx import ASGITransport, AsyncClient
from openpyxl import load_workbook
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker, create_async_engine

from backend.app import create_app
from backend.database import Base, get_db

# Test database - in-memory SQLite
TEST_DATABASE_URL = "sqlite+aiosqlite:///:memory:"
test_engine = create_async_engine(TEST_DATABASE_URL, echo=False)
TestSessionLocal = async_sessionmaker(test_engine, class_=AsyncSession, expire_on_commit=False)

AUTH_HEADERS = {"Authorization": "Bearer change-me-in-production"}


async def override_get_db():
    async with TestSessionLocal() as session:
        yield session


@pytest_asyncio.fixture
async def app():
    """Create test app with test database."""
    application = create_app()
    application.dependency_overrides[get_db] = override_get_db

    async with test_engine.begin() as conn:
        await conn.run_sync(Base.metadata.create_all)

    yield application

    async with test_engine.begin() as conn:
        await conn.run_sync(Base.metadata.drop_all)


@pytest_asyncio.fixture
async def client(app):
    """Create async test client."""
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as ac:
        yield ac


@pytest.mark.asyncio
async def test_export_t13_returns_xlsx(client):
    """Test that endpoint returns a valid xlsx file."""
    response = await client.get(
        "/api/v1/timesheet/export-t13",
        params={"year": 2025, "month": 6},
        headers=AUTH_HEADERS,
    )
    assert response.status_code == 200
    assert (
        response.headers["content-type"]
        == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert "timesheet_T13_2025_06.xlsx" in response.headers["content-disposition"]

    # Verify it's a valid workbook
    wb = load_workbook(io.BytesIO(response.content))
    assert wb.active is not None


@pytest.mark.asyncio
async def test_export_t13_requires_auth(client):
    """Test that endpoint requires authentication."""
    response = await client.get(
        "/api/v1/timesheet/export-t13",
        params={"year": 2025, "month": 6},
    )
    assert response.status_code == 401


@pytest.mark.asyncio
async def test_export_t13_with_data(client):
    """Test export with employees and marks populated."""
    # Create employees
    resp = await client.post(
        "/api/v1/employees",
        json={"fio": "Иванова А.Б.", "position": "Воспитатель", "rate": 1.0},
        headers=AUTH_HEADERS,
    )
    assert resp.status_code == 201
    emp1_id = resp.json()["id"]

    resp = await client.post(
        "/api/v1/employees",
        json={"fio": "Петрова В.Г.", "position": "Няня", "rate": 0.5},
        headers=AUTH_HEADERS,
    )
    assert resp.status_code == 201
    emp2_id = resp.json()["id"]

    # Add marks
    await client.post(
        "/api/v1/timesheet/mark",
        json={"employee_id": emp1_id, "date": "2025-06-01", "mark_type": "Я"},
        headers=AUTH_HEADERS,
    )
    await client.post(
        "/api/v1/timesheet/mark",
        json={"employee_id": emp1_id, "date": "2025-06-02", "mark_type": "Я"},
        headers=AUTH_HEADERS,
    )
    await client.post(
        "/api/v1/timesheet/mark",
        json={"employee_id": emp1_id, "date": "2025-06-03", "mark_type": "Б"},
        headers=AUTH_HEADERS,
    )
    await client.post(
        "/api/v1/timesheet/mark",
        json={"employee_id": emp2_id, "date": "2025-06-01", "mark_type": "Я"},
        headers=AUTH_HEADERS,
    )

    # Export
    response = await client.get(
        "/api/v1/timesheet/export-t13",
        params={"year": 2025, "month": 6},
        headers=AUTH_HEADERS,
    )
    assert response.status_code == 200

    # Verify workbook structure
    wb = load_workbook(io.BytesIO(response.content))
    ws = wb.active
    assert ws.title == "Табель Т-13"

    # Check header
    assert ws["A1"].value == "Детский сад"
    assert "06/2025" in ws["A3"].value

    # Header row is row 5
    assert ws.cell(row=5, column=1).value == "No п/п"
    assert ws.cell(row=5, column=2).value == "ФИО"

    # Employee data starts at row 6
    assert ws.cell(row=6, column=2).value == "Иванова А.Б."
    assert ws.cell(row=7, column=2).value == "Петрова В.Г."

    # Check day marks for first employee (day 1 = col 5, day 2 = col 6, day 3 = col 7)
    assert ws.cell(row=6, column=5).value == "Я"   # day 1
    assert ws.cell(row=6, column=6).value == "Я"   # day 2
    assert ws.cell(row=6, column=7).value == "Б"   # day 3

    # Check second employee day 1
    assert ws.cell(row=7, column=5).value == "Я"   # day 1

    # June has 30 days, so totals start at col 5 + 30 = 35
    totals_col = 5 + 30  # col 35
    # Employee 1: total days = 3
    assert ws.cell(row=6, column=totals_col).value == 3
    # Employee 2: total days = 1
    assert ws.cell(row=7, column=totals_col).value == 1

    # Bottom totals row (row 8 = header_row + 2 employees + 1)
    assert ws.cell(row=8, column=2).value == "ИТОГО"
