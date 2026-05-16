"""OCR document scanning endpoints."""

import io
import json
import logging

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font, Border, Side, Alignment

from backend.auth import verify_bearer_token
from backend.ai.groq_client import vision_completion
from backend.schemas.ocr import OcrProcessRequest, OcrProcessResponse, OcrToExcelRequest

logger = logging.getLogger(__name__)

router = APIRouter(prefix="/ocr", tags=["ocr"])

OCR_PROMPT = (
    "Extract all numbers, dates, names, amounts, INN from this Russian accounting document. "
    "Return a JSON object with doc_type (invoice/timesheet/receipt/payslip) and fields dict "
    "mapping field names to values. Field names should be in Russian."
)


@router.post("/process", response_model=OcrProcessResponse)
async def ocr_process(
    data: OcrProcessRequest,
    _token: str = Depends(verify_bearer_token),
):
    """Process an image through OCR using Groq Vision API."""
    if not data.image_base64:
        raise HTTPException(status_code=400, detail="image_base64 is required")

    result = await vision_completion(data.image_base64, OCR_PROMPT)

    if not result:
        raise HTTPException(status_code=502, detail="AI service unavailable")

    # Try to parse JSON from the LLM response
    try:
        # The LLM may wrap JSON in markdown code blocks
        clean = result.strip()
        if clean.startswith("```"):
            # Remove markdown code block markers
            lines = clean.split("\n")
            lines = [ln for ln in lines if not ln.strip().startswith("```")]
            clean = "\n".join(lines)
        parsed = json.loads(clean)
    except (json.JSONDecodeError, ValueError):
        # If parsing fails, return the raw text as unknown document
        return OcrProcessResponse(
            doc_type="unknown",
            fields={},
            raw_text=result,
        )

    doc_type = parsed.get("doc_type", "unknown")
    fields = parsed.get("fields", {})
    # Ensure fields values are strings
    fields = {str(k): str(v) for k, v in fields.items()}

    return OcrProcessResponse(
        doc_type=doc_type,
        fields=fields,
        raw_text=result,
    )


@router.post("/to-excel")
async def ocr_to_excel(
    data: OcrToExcelRequest,
    _token: str = Depends(verify_bearer_token),
):
    """Convert structured OCR data to Excel file."""
    wb = Workbook()
    ws = wb.active
    ws.title = data.title[:31]  # Excel sheet name max 31 chars

    # Styles
    header_font = Font(bold=True, size=12)
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    # Title row
    ws.merge_cells("A1:B1")
    title_cell = ws["A1"]
    title_cell.value = f"{data.title} ({data.doc_type})"
    title_cell.font = Font(bold=True, size=14)
    title_cell.alignment = Alignment(horizontal="center")

    # Column widths
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 40

    # Headers
    ws["A3"] = "\u041f\u043e\u043b\u0435"
    ws["B3"] = "\u0417\u043d\u0430\u0447\u0435\u043d\u0438\u0435"
    ws["A3"].font = header_font
    ws["B3"].font = header_font
    ws["A3"].border = thin_border
    ws["B3"].border = thin_border

    # Data rows
    row = 4
    for key, value in data.fields.items():
        ws.cell(row=row, column=1, value=key).border = thin_border
        ws.cell(row=row, column=2, value=value).border = thin_border
        row += 1

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    filename = f"ocr_{data.doc_type}.xlsx"

    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )
