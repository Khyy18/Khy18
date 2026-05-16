"""Timesheet endpoints."""

import calendar
import io

from fastapi import APIRouter, Depends, Query
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import select, func
from sqlalchemy.ext.asyncio import AsyncSession

from backend.auth import verify_bearer_token
from backend.database import get_db
from backend.models.employee import Employee
from backend.models.timesheet import TimesheetMark
from backend.schemas.timesheet import TimesheetMarkCreate, TimesheetMarkResponse
from backend.security.encryption import decrypt_value

router = APIRouter(prefix="/timesheet", tags=["timesheet"])


@router.post("/mark", response_model=TimesheetMarkResponse, status_code=201)
async def add_mark(
    data: TimesheetMarkCreate,
    db: AsyncSession = Depends(get_db),
    _token: str = Depends(verify_bearer_token),
):
    mark = TimesheetMark(
        employee_id=data.employee_id,
        date=data.date,
        mark_type=data.mark_type,
    )
    db.add(mark)
    await db.commit()
    await db.refresh(mark)
    return mark


@router.get("/summary/{employee_id}")
async def timesheet_summary(
    employee_id: int,
    year: int = Query(...),
    month: int = Query(...),
    db: AsyncSession = Depends(get_db),
):
    """Monthly summary: count marks by type for given employee/month."""
    date_prefix = f"{year:04d}-{month:02d}"
    result = await db.execute(
        select(TimesheetMark.mark_type, func.count(TimesheetMark.id))
        .where(
            TimesheetMark.employee_id == employee_id,
            TimesheetMark.date.like(f"{date_prefix}%"),
        )
        .group_by(TimesheetMark.mark_type)
    )
    rows = result.all()
    summary = {row[0]: row[1] for row in rows}
    return {"employee_id": employee_id, "year": year, "month": month, "summary": summary}


# Mark type codes used in T-13 form
MARK_CODES = ["Я", "Б", "О", "ОТ", "П"]


@router.get("/export-t13")
async def export_t13(
    year: int = Query(...),
    month: int = Query(...),
    db: AsyncSession = Depends(get_db),
    _token: str = Depends(verify_bearer_token),
):
    """Export T-13 timesheet form as Excel file."""
    days_in_month = calendar.monthrange(year, month)[1]

    # Fetch all employees
    emp_result = await db.execute(select(Employee).order_by(Employee.id))
    employees = emp_result.scalars().all()

    # Fetch all marks for the month
    date_prefix = f"{year:04d}-{month:02d}"
    marks_result = await db.execute(
        select(TimesheetMark).where(TimesheetMark.date.like(f"{date_prefix}%"))
    )
    marks = marks_result.scalars().all()

    # Build lookup: {employee_id: {day_number: mark_type}}
    marks_lookup: dict[int, dict[int, str]] = {}
    for mark in marks:
        day = int(mark.date.split("-")[2])
        marks_lookup.setdefault(mark.employee_id, {})[day] = mark.mark_type

    # Create workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Табель Т-13"

    # Styles
    header_font = Font(name="Times New Roman", size=12, bold=True)
    data_font = Font(name="Times New Roman", size=10)
    bold_font = Font(name="Times New Roman", size=10, bold=True)
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )
    center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # Header rows
    ws.merge_cells("A1:F1")
    ws["A1"] = "Детский сад"
    ws["A1"].font = header_font

    ws.merge_cells("A2:F2")
    ws["A2"] = "Структурное подразделение"
    ws["A2"].font = data_font

    ws.merge_cells("A3:F3")
    ws["A3"] = f"Табель учёта рабочего времени за {month:02d}/{year}"
    ws["A3"].font = header_font

    # Table header row (row 5)
    header_row = 5
    headers = ["No п/п", "ФИО", "Должность", "Таб. номер"]
    for col_idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=col_idx, value=h)
        cell.font = bold_font
        cell.border = thin_border
        cell.alignment = center_align

    # Day columns
    day_start_col = len(headers) + 1
    for day in range(1, days_in_month + 1):
        cell = ws.cell(row=header_row, column=day_start_col + day - 1, value=day)
        cell.font = bold_font
        cell.border = thin_border
        cell.alignment = center_align

    # Totals columns after days
    totals_start_col = day_start_col + days_in_month
    total_headers = ["Итого дней"] + [f"Итого {code}" for code in MARK_CODES]
    for i, th in enumerate(total_headers):
        cell = ws.cell(row=header_row, column=totals_start_col + i, value=th)
        cell.font = bold_font
        cell.border = thin_border
        cell.alignment = center_align

    # Set column widths
    ws.column_dimensions["A"].width = 6
    ws.column_dimensions["B"].width = 25
    ws.column_dimensions["C"].width = 20
    ws.column_dimensions["D"].width = 10
    for day in range(1, days_in_month + 1):
        col_letter = get_column_letter(day_start_col + day - 1)
        ws.column_dimensions[col_letter].width = 4

    # Employee data rows
    # Track totals per mark type across all employees
    grand_totals: dict[str, int] = {code: 0 for code in MARK_CODES}

    for row_idx, emp in enumerate(employees, start=header_row + 1):
        ws.cell(row=row_idx, column=1, value=row_idx - header_row).font = data_font
        ws.cell(row=row_idx, column=1).border = thin_border
        ws.cell(row=row_idx, column=1).alignment = center_align

        ws.cell(row=row_idx, column=2, value=decrypt_value(emp.fio)).font = data_font
        ws.cell(row=row_idx, column=2).border = thin_border

        ws.cell(row=row_idx, column=3, value=emp.position).font = data_font
        ws.cell(row=row_idx, column=3).border = thin_border

        ws.cell(row=row_idx, column=4, value=emp.id).font = data_font
        ws.cell(row=row_idx, column=4).border = thin_border
        ws.cell(row=row_idx, column=4).alignment = center_align

        emp_marks = marks_lookup.get(emp.id, {})
        total_days = 0
        emp_mark_counts: dict[str, int] = {code: 0 for code in MARK_CODES}

        for day in range(1, days_in_month + 1):
            col = day_start_col + day - 1
            mark_code = emp_marks.get(day, "")
            cell = ws.cell(row=row_idx, column=col, value=mark_code)
            cell.font = data_font
            cell.border = thin_border
            cell.alignment = center_align
            if mark_code:
                total_days += 1
                if mark_code in emp_mark_counts:
                    emp_mark_counts[mark_code] += 1

        # Totals for this employee
        cell = ws.cell(row=row_idx, column=totals_start_col, value=total_days)
        cell.font = data_font
        cell.border = thin_border
        cell.alignment = center_align

        for i, code in enumerate(MARK_CODES):
            count = emp_mark_counts[code]
            cell = ws.cell(row=row_idx, column=totals_start_col + 1 + i, value=count)
            cell.font = data_font
            cell.border = thin_border
            cell.alignment = center_align
            grand_totals[code] += count

    # Bottom totals row
    totals_row = header_row + len(employees) + 1
    ws.cell(row=totals_row, column=1, value="").border = thin_border
    ws.cell(row=totals_row, column=2, value="ИТОГО").font = bold_font
    ws.cell(row=totals_row, column=2).border = thin_border
    ws.cell(row=totals_row, column=3, value="").border = thin_border
    ws.cell(row=totals_row, column=4, value="").border = thin_border

    for day in range(1, days_in_month + 1):
        ws.cell(row=totals_row, column=day_start_col + day - 1, value="").border = thin_border

    grand_total_days = sum(grand_totals.values())
    cell = ws.cell(row=totals_row, column=totals_start_col, value=grand_total_days)
    cell.font = bold_font
    cell.border = thin_border
    cell.alignment = center_align

    for i, code in enumerate(MARK_CODES):
        cell = ws.cell(row=totals_row, column=totals_start_col + 1 + i, value=grand_totals[code])
        cell.font = bold_font
        cell.border = thin_border
        cell.alignment = center_align

    # Save to buffer
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    filename = f"timesheet_T13_{year}_{month:02d}.xlsx"
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )
